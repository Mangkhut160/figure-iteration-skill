#!/usr/bin/env python3
"""List the shapes of a pptx slide (read-only): what a patch can find them by, and where they are.

usage: python pptx_inspect.py <pptx> [--group 'n:<id> '] [--assets DIR]
Without --group: one line per top-level shape (kind, name, x, y, w, h in pt; for a group the texts inside it and [moved] when
the user has moved or resized it; for any other shape its text or picture asset and its fill colour). This is the view of a
deck without groups (drawn by hand in an editor).
With --group: one line per shape inside the groups whose name starts with that prefix: kind, text or picture asset,
fill colour, x, y, w, h and the bottom edge (y + h) in pt. These coordinates are in the group's child space; when the group was
moved the header says by how much (page position = child position + that shift; pptx_kit.page_geom does it).
Use it before writing a patch (names, colours, exact texts) and after it (read back what the patch set, e.g. bar heights).
Fill colours: a hex value is an explicit RGB fill (a patch can match it); theme:<name> is a palette colour set on the shape;
style:<name> means the shape has no fill of its own and the theme's style paints it; none = no fill.
"""
import hashlib
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation  # noqa: E402

from _deck import E, fill_desc, group_shift, is_identity, kind as _kind, known_assets  # noqa: E402


def kind(sh):
    return _kind(sh)[:10]


def label(sh, known):
    if sh.shape_type == 13:
        return known.get(hashlib.sha1(sh.image.blob).hexdigest(), 'picture (not in assets/)')
    return sh.text_frame.text.strip().replace('\n', ' / ') if sh.has_text_frame else ''


def row(sh, known, indent=''):
    x, y, w, h = (v / E for v in (sh.left, sh.top, sh.width, sh.height))
    return (f'{indent}{kind(sh):<10} {label(sh, known)[:34]!r:<36} {fill_desc(sh):<13} '
            f'x {x:7.2f}  y {y:7.2f}  w {w:7.2f}  h {h:7.2f}  bottom {y + h:7.2f}')


def shift_note(g):
    if is_identity(g):
        return ''
    dx, dy, sx, sy = group_shift(g)
    s = f'child space is shifted by ({dx / E:+.2f}, {dy / E:+.2f}) pt on the page'
    return s + (f', scaled x{sx:.3f} / x{sy:.3f}' if abs(sx - 1) > 1e-4 or abs(sy - 1) > 1e-4 else '')


def walk(shapes, known, indent):
    for c in shapes:
        print(row(c, known, indent))
        if c.shape_type == 6:
            note = shift_note(c)
            if note:
                print(f'{indent}  ({note})')
            walk(c.shapes, known, indent + '  ')


def main(path, prefix, assets):
    known = known_assets(assets)
    slide = Presentation(path).slides[0]
    if prefix is None:
        for s in slide.shapes:
            if s.shape_type == 6:
                extra = '; '.join(c.text_frame.text.strip() for c in s.shapes if c.has_text_frame and c.text_frame.text.strip())[:60]
                if not is_identity(s):
                    extra = '[moved] ' + extra
            else:
                extra = f'{label(s, known)[:34]!r} {fill_desc(s)}'.strip()
            x, y, w, h = (v / E for v in (s.left, s.top, s.width, s.height))
            print(f'{kind(s):<10} {s.name[:42]:<42} x {x:6.1f} y {y:6.1f} w {w:6.1f} h {h:5.1f}  {extra}')
        return
    groups = [s for s in slide.shapes if s.name.startswith(prefix)]
    if not groups:
        raise SystemExit(f'no top-level shape named {prefix!r}*')
    for g in groups:
        print(f'== {g.name}')
        if g.shape_type == 6:
            note = shift_note(g)
            if note:
                print(f'   ({note})')
            walk(g.shapes, known, '  ')
        else:
            print(row(g, known, '  '))


if __name__ == '__main__':
    args = sys.argv[1:]
    opts = {'--group': None, '--assets': os.path.join(os.environ['FIG_DIR'], 'assets') if os.environ.get('FIG_DIR') else None}
    for flag in list(opts):
        if flag in args:
            i = args.index(flag)
            opts[flag] = args[i + 1]
            del args[i:i + 2]
    if len(args) != 1:
        raise SystemExit(__doc__)
    main(args[0], opts['--group'], opts['--assets'])
