#!/usr/bin/env python3
"""A small test deck for dry-running a new paper's setup (python with python-pptx).

  python make_test_deck.py <out.pptx> --page <WIDTH>x<HEIGHT> --size <pt>
Three boxes, two arrows, a label and a small bitmap on a page of the given size in pt (the venue's figure width x a height), all
text in Arial at --size. Feed it to make_from_saved.sh with templates/patch_template.py (a patch that changes nothing) and the
paper's figure.env to see, before any real figure exists, that rendering, the page-size and font-size checks and the metadata
checks work with this paper's numbers. A hand-drawn deck is flat (no groups, no figgenie names), like this one.
"""
import os
import sys
import tempfile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
args = sys.argv[1:]
if not args:
    raise SystemExit(__doc__)
out = args[0]
page, size = None, None
for flag in ('--page', '--size'):
    if flag in args:
        v = args[args.index(flag) + 1]
        page, size = (v, size) if flag == '--page' else (page, float(v))
if page is None or size is None:
    raise SystemExit(__doc__)
W, H = (float(v) for v in page.lower().split('x'))
if os.path.exists(out):
    raise SystemExit(f'{out} exists: pick a new name')

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def style(tf, align=PP_ALIGN.CENTER):
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size, r.font.name = Pt(size), 'Arial'
            r.font.color.rgb = RGBColor(0, 0, 0)


bw, bh, gap = W * 0.24, H * 0.26, W * 0.1
x0, y0 = (W - 3 * bw - 2 * gap) / 2, H * 0.3
for i, (text, rgb) in enumerate((('Input', 'D9D9D9'), ('Module', 'BDD7EE'), ('Output', 'A9D18E'))):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x0 + i * (bw + gap)), Pt(y0), Pt(bw), Pt(bh))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(rgb)
    sh.line.color.rgb = RGBColor(0x59, 0x59, 0x59)
    sh.line.width = Pt(0.5)
    sh.text_frame.text = text
    style(sh.text_frame)
for i in range(2):
    xa = x0 + i * (bw + gap) + bw
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(xa), Pt(y0 + bh / 2), Pt(xa + gap), Pt(y0 + bh / 2))
    cn.line.width = Pt(0.9)
    cn.line.color.rgb = RGBColor(0x26, 0x26, 0x26)
    etree.SubElement(cn._element.spPr.find(A + 'ln'), A + 'tailEnd').set('type', 'triangle')
tb = slide.shapes.add_textbox(Pt(x0), Pt(H * 0.78), Pt(bw * 2), Pt(size * 2))
tb.text_frame.text = 'legend text'
style(tb.text_frame, PP_ALIGN.LEFT)
with tempfile.TemporaryDirectory() as td:
    png = os.path.join(td, 'dot.png')
    Image.new('RGB', (200, 200), (230, 120, 40)).save(png)                 # 200 px on 12 pt = 1200 ppi
    slide.shapes.add_picture(png, Pt(W - 20), Pt(H * 0.74), Pt(12), Pt(12))
    cp = prs.core_properties
    cp.author = cp.last_modified_by = cp.comments = cp.title = ''
    prs.save(out)
print('wrote', out, f'({W:g} x {H:g} pt, text {size:g} pt)')
