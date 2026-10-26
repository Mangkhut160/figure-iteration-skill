"""Helpers for the patch scripts that change a copy of the pptx the user saved (python-pptx + lxml).

Shapes are found by group name, text, fill colour or picture bytes, never by position, so the user's own moves survive
when the patch is re-run on a newer save. Geometry is in pt (1 pt = 12700 EMU). In decks written by figgenie's svg2pptx
the group child transforms are identity (chOff = off, chExt = ext), so child coordinates are slide coordinates.
FIG_DIR (environment, set by make_from_saved.sh) is the figure folder: pictures are named by their bytes against
FIG_DIR/assets/. K.formula() typesets TeX with the local LaTeX (latex_math.py) and inserts the glyphs as freeform shapes.
A deck without groups (drawn by hand in an editor): pass K.slide where the finders ask for a group (K.text_of(K.slide, 'Head'))
and K.slide or None to the adders (clone, text_box, add_picture, formula); the shapes are then found by text, explicit RGB
fill or picture bytes. pptx_inspect.py shows what a patch can match.
"""
import copy
import hashlib
import os
import re
import subprocess
import sys
import tempfile
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import _paths  # noqa: E402
from _deck import E, A, page_geom  # noqa: E402,F401

FG = _paths.figgenie_scripts()          # None when figgenie is not installed: only tw() and formula() need it and stop with a message


def _need_figgenie():
    if not FG:
        raise SystemExit('figgenie-paper-diagram/scripts not found: set FIGGENIE_SCRIPTS (tw() and formula() need it)')
    if FG not in sys.path:
        sys.path.insert(0, FG)


CTR_K = 0.3538                  # baseline = box height / 2 + CTR_K * size in a middle-anchored box (svg2pptx text boxes are 2.2 x size high; checked at 6, 6.5 and 7 pt: error < 0.1 pt)
INK = os.environ.get('FIG_INK', '000000')       # text colour of new formulas; set FIG_INK in figure.env for another ink
FIG_DIR = os.environ.get('FIG_DIR') or os.getcwd()
ASSETS = os.path.join(FIG_DIR, 'assets')


def tw(s, size, bold=False, font=None):
    """Advance width of `s` in pt (figgenie's font metrics)."""
    _need_figgenie()
    from measure_text import width_of
    return width_of(s, font or os.environ.get('FIG_FONT', 'arial'), size, bold)[0]


def pt(v):
    return v / E


def geom(sh):
    return pt(sh.left), pt(sh.top), pt(sh.width), pt(sh.height)


def put(sh, x=None, y=None, w=None, h=None):
    for attr, v in (('left', x), ('top', y), ('width', w), ('height', h)):
        if v is not None:
            setattr(sh, attr, Emu(round(v * E)))


def shift(sh, dx=0.0, dy=0.0):
    x, y, _, _ = geom(sh)
    put(sh, x + dx, y + dy)


def fill_of(sh):
    v = sh._element.xpath('./p:spPr/a:solidFill/a:srgbClr/@val')
    return v[0].upper() if v else None


def set_fill(sh, rgb):
    """Solid fill `rgb` (hex). An explicit RGB fill keeps its transparency; a theme colour or a fill painted by the theme
    style is replaced by an explicit one."""
    hit = sh._element.xpath('./p:spPr/a:solidFill/a:srgbClr')
    if hit:
        hit[0].set('val', rgb)
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(rgb.upper())


def set_line(sh, rgb, width_pt):
    hit = sh._element.xpath('./p:spPr/a:ln/a:solidFill/a:srgbClr')
    if hit:
        sh._element.xpath('./p:spPr/a:ln')[0].set('w', str(round(width_pt * E)))
        hit[0].set('val', rgb)
    else:
        sh.line.color.rgb = RGBColor.from_string(rgb.upper())
        sh.line.width = Emu(round(width_pt * E))


def set_text(sh, text):
    """New text in the first run of the first paragraph (keeps its font, size, colour); other runs are dropped."""
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def put_text_in_cell(cell, text, color, ref_run):
    """Write `text` into a box shape with zero insets, styled like `ref_run` (an a:r element of an existing label)."""
    body = cell._element.txBody
    for k in ('lIns', 'rIns', 'tIns', 'bIns'):
        body.bodyPr.set(k, '0')
    p = body.xpath('./a:p')[0]
    for old in p.xpath('./a:r'):
        p.remove(old)
    run = copy.deepcopy(ref_run)
    run.xpath('./a:t')[0].text = text
    run.xpath('./a:rPr/a:solidFill/a:srgbClr')[0].set('val', color)
    p.append(run)


def remove(sh):
    sh._element.getparent().remove(sh._element)


def one(items, what):
    """The only element of `items` (a list from pics / boxes / glyphs / texts), or stop with a message."""
    items = list(items)
    if len(items) != 1:
        raise SystemExit(f'expected one {what}, found {len(items)}')
    return items[0]


def _asset_hashes():
    out = {}
    if os.path.isdir(ASSETS):
        for root, _, files in os.walk(ASSETS):
            for fn in files:
                if fn.lower().endswith(('.png', '.jpg', '.jpeg')):
                    p = os.path.join(root, fn)
                    with open(p, 'rb') as fh:
                        out[hashlib.sha1(fh.read()).hexdigest()] = os.path.relpath(p, ASSETS).replace(os.sep, '/')
    return out


KNOWN = _asset_hashes()


def asset_of(pic):
    """assets/-relative name of a picture shape, or None when its bytes match no file in FIG_DIR/assets/."""
    return KNOWN.get(hashlib.sha1(pic.image.blob).hexdigest())


class Kit:
    """K = Kit(Presentation(copy_of_saved)); find and change shapes; K.save(new_path) at the end."""

    def __init__(self, prs, slide_index=0):
        self.prs = prs
        self.slide = prs.slides[slide_index]
        self.touched = []                  # groups whose children changed: extents recalculated in save()

    # ------------------------------------------------------------------ finding (exactly one, or stop with a message)
    def group(self, prefix):
        """The one top-level group whose name starts with `prefix`. svg2pptx names a group '<spec id> <label>', so pass
        the id with a trailing space ('n:enc ') to avoid matching 'n:encoder'."""
        g = [s for s in self.slide.shapes if s.name.startswith(prefix) and s.shape_type == 6]
        if len(g) != 1:
            raise SystemExit(f'expected one group named {prefix!r}, found {len(g)}')
        self.touched.append(g[0])
        return g[0]

    def top(self, prefix):
        s = [s for s in self.slide.shapes if s.name.startswith(prefix) and s.shape_type != 6]
        if len(s) != 1:
            raise SystemExit(f'expected one top-level shape named {prefix!r}, found {len(s)}')
        return s[0]

    @staticmethod
    def text_of(g, text):
        s = [c for c in g.shapes if c.has_text_frame and c.text_frame.text.strip() == text]
        if len(s) != 1:
            raise SystemExit(f'expected one text {text!r} in {g.name!r}, found {len(s)}; several equal labels: use texts()')
        return s[0]

    @staticmethod
    def texts(g, text):
        """All shapes of `g` whose text is `text`, left to right. Tell equal labels apart by what they belong to (e.g. the
        bar whose centre x matches), not by absolute position."""
        return sorted((c for c in g.shapes if c.has_text_frame and c.text_frame.text.strip() == text), key=lambda c: c.left)

    @staticmethod
    def pics(g, asset):
        """Pictures in `g` whose bytes are FIG_DIR/assets/<asset> (e.g. 'icons/a.png')."""
        return [c for c in g.shapes if c.shape_type == 13 and asset_of(c) == asset]

    @staticmethod
    def boxes(g, fill=None):
        return [c for c in g.shapes if c.shape_type == 1 and (fill is None or fill_of(c) == fill.upper())]

    @staticmethod
    def biggest_box(g):
        return max((c for c in g.shapes if c.shape_type == 1), key=lambda c: c.width * c.height)

    @staticmethod
    def glyphs(g):
        """Freeform shapes of `g` left to right (LaTeX glyphs are freeforms in these decks)."""
        return sorted((c for c in g.shapes if c.shape_type == 5), key=lambda c: c.left)

    # ------------------------------------------------------------------ adding
    def new_id(self):
        return str(max(int(i) for i in self.slide.shapes._spTree.xpath('.//p:cNvPr/@id')) + 1)

    def _dst(self, g):
        """The container for adding: a group, or None for the slide (K.slide is accepted as the slide)."""
        return None if g is None or g is self.slide else g

    def _proxy(self, dst_group, el):
        shapes = dst_group.shapes if dst_group is not None else self.slide.shapes
        return next(c for c in shapes if c._element is el)

    def clone(self, dst_group, src, name=None):
        """Copy `src` (any shape, any group) to the end of `dst_group` (None or K.slide = the slide); returns the new shape."""
        dst_group = self._dst(dst_group)
        el = copy.deepcopy(src._element)
        (dst_group._element if dst_group is not None else self.slide.shapes._spTree).append(el)
        el.xpath('.//p:cNvPr')[0].set('id', self.new_id())
        if name:
            el.xpath('.//p:cNvPr')[0].set('name', name)
        return self._proxy(dst_group, el)

    def text_box(self, dst_group, template, text, x, baseline, w=None, top=None):
        """A one-line label copied from `template` starting at x, with the text on `baseline` (pt). The box top is worked out
        from the template's own height and font size, which is right for a middle-anchored box (what svg2pptx writes); for a
        box anchored any other way pass `top` (the box's top edge in pt) instead."""
        sh = self.clone(dst_group, template)
        set_text(sh, text)
        size = float((template._element.xpath('.//a:rPr/@sz') or ['1800'])[0]) / 100        # no explicit size: the 18 pt default
        bold = bool(template._element.xpath('.//a:rPr[@b="1"]'))
        if top is None:
            if (template._element.xpath('./p:txBody/a:bodyPr/@anchor') or ['t'])[0] != 'ctr':
                raise SystemExit('text_box: the template is not middle-anchored: pass top= (the box top edge in pt)')
            top = baseline - (pt(template.height) / 2 + CTR_K * size)
        put(sh, x=x, y=top, w=w if w else tw(text, size, bold) + 1.0)
        return sh

    def add_picture(self, g, png, x, y, h, name, descr='', min_ppi=None):
        """A picture `h` pt high at (x, y), its own aspect ratio; refuses when it would print below `min_ppi` (default: FIG_MIN_PPI, else 300)."""
        from PIL import Image
        g = self._dst(g)
        min_ppi = float(min_ppi or os.environ.get('FIG_MIN_PPI', 300))
        with Image.open(png) as im:
            pw, ph = im.size
        w = h * pw / ph
        ppi = pw / (w / 72)
        if ppi < min_ppi:
            raise SystemExit(f'{png}: {ppi:.0f} ppi at {w:.1f} pt wide (< {min_ppi}); use a larger source')
        pic = (g.shapes if g is not None else self.slide.shapes).add_picture(
            png, Emu(round(x * E)), Emu(round(y * E)), Emu(round(w * E)), Emu(round(h * E)))
        pic._element.nvPicPr.cNvPr.set('id', self.new_id())     # python-pptx counts ids per group only
        pic.name = name
        if descr:
            pic._element.nvPicPr.cNvPr.set('descr', descr)
        return pic

    def formula(self, dst_group, tex, x, baseline, size, color=None, name=None):
        """TeX (math mode, without $) typeset by the local LaTeX and inserted as freeform glyphs: the box starts at x, its
        baseline at `baseline` (pt). Returns (shapes, (w, h, d)) = box width, height above and depth below the baseline.
        Colour: `color` (hex), else FIG_INK, else 000000. The math size is the profile's choice (a size up from the label
        beside it when the math font's x-height is lower: Computer Modern next to Arial)."""
        from latex_math import typeset
        _need_figgenie()
        color = color or INK
        dst_group = self._dst(dst_group)
        v = typeset([(tex, size)])[(tex, size)]
        W, H = max(v['w'] + 2, 80), max(v['h'] + v['d'] + 2, 80)       # a pptx page must be at least 1 inch
        svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W:.4f} {H:.4f}" width="{W:.4f}pt" height="{H:.4f}pt">'
               f'<g id="tex" transform="translate(1 {1 + v["h"]:.4f})" fill="#{color}">{v["body"]}</g></svg>')
        with tempfile.TemporaryDirectory() as td:
            src, out = os.path.join(td, 'f.svg'), os.path.join(td, 'f.pptx')
            with open(src, 'w', encoding='utf-8') as fh:
                fh.write(svg)
            r = subprocess.run([sys.executable, os.path.join(FG, 'svg2pptx.py'), src, '-o', out], capture_output=True, text=True)
            if r.returncode or not os.path.exists(out):
                raise SystemExit(f'svg2pptx failed on {tex!r}:\n{r.stderr[-1500:]}')
            glyph_els = Presentation(out).slides[0].shapes._spTree.xpath('.//p:sp')
            dx, dy = round((x - 1) * E), round((baseline - 1 - v['h']) * E)
            parent = dst_group._element if dst_group is not None else self.slide.shapes._spTree
            new = []
            for sp in glyph_els:
                el = copy.deepcopy(sp)
                for st in el.xpath('./p:style'):              # theme effect reference: LibreOffice draws it as a shadow
                    el.remove(st)
                off = el.xpath('./p:spPr/a:xfrm/a:off')[0]
                off.set('x', str(int(off.get('x')) + dx))
                off.set('y', str(int(off.get('y')) + dy))
                parent.append(el)
                c = el.xpath('./p:nvSpPr/p:cNvPr')[0]
                c.set('id', self.new_id())
                c.set('name', f'{name or "tex"} {tex}')
                c.set('descr', f'baseline={baseline:.4f}; size={size:g}')     # replace_formula reads the baseline back
                new.append(self._proxy(dst_group, el))
        return new, (v['w'], v['h'], v['d'])

    def replace_formula(self, g, old, tex, size, align='left', color=None, name=None):
        """Remove the glyph shapes `old` (one whole formula) and typeset `tex` in their place: same baseline, and the ink's
        left edge (align='left') or right edge (align='right') where the old ink was. The baseline comes from the descr
        K.formula writes; for older glyphs it is the bottom of the first glyph (exact unless that glyph has a descender,
        e.g. '[' or 'p': then pass glyphs that start with a plain letter, or set the baseline yourself with formula())."""
        old = sorted(old, key=lambda c: c.left)
        if not old:
            raise SystemExit('replace_formula: no glyphs given')
        descr = (old[0]._element.xpath('./p:nvSpPr/p:cNvPr/@descr') or [''])[0]
        meta = dict(kv.split('=', 1) for kv in descr.split('; ') if '=' in kv)
        base = float(meta['baseline']) if 'baseline' in meta else pt(old[0].top + old[0].height)
        left, right = pt(min(c.left for c in old)), pt(max(c.left + c.width for c in old))
        for c in old:
            remove(c)
        new, box = self.formula(g, tex, left, base, size, color, name)
        ink_l, ink_r = pt(min(c.left for c in new)), pt(max(c.left + c.width for c in new))
        dx = (left - ink_l) if align == 'left' else (right - ink_r)
        for c in new:
            shift(c, dx)
        return new, box

    # ------------------------------------------------------------------ writing
    def save(self, outp):
        """Refit the child space of the groups whose children changed (a group nobody touched is left exactly as the user
        saved it, and a group he moved stays where he put it), blank the author fields and the Company / Manager of app.xml
        (editors write the account and organisation names there), blank a document comment that names a tool, refuse
        duplicate shape ids and an existing output file, then write."""
        if os.path.exists(outp):
            raise SystemExit(f'{outp} exists: pick a new name (never overwrite a file the user may have open)')
        for g in self.touched:
            _refit(g._element)
        cp = self.prs.core_properties
        cp.author = ''
        cp.last_modified_by = ''
        if re.search(r'generated|python-pptx|created (with|by)', cp.comments or '', re.I):
            cp.comments = ''
        ids = [int(i) for i in self.slide.shapes._spTree.xpath('.//p:cNvPr/@id')]
        dup = sorted({i for i in ids if ids.count(i) > 1})
        if dup:
            raise SystemExit(f'duplicate shape ids {dup}')
        self.prs.save(outp)
        _blank_app_props(outp)
        print('wrote', outp)


def _child_rect(el):
    """(x, y, cx, cy) in EMU of a child element in its parent's space, or None."""
    for path in ('./p:spPr/a:xfrm', './p:grpSpPr/a:xfrm', './p:xfrm'):
        x = el.xpath(path)
        if x and x[0].find(A + 'off') is not None and x[0].find(A + 'ext') is not None:
            o, e = x[0].find(A + 'off'), x[0].find(A + 'ext')
            return int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy'))
    return None


def _refit(grp):
    """Fit the group's child space to its children without moving anything on the page: the group keeps its own position
    and scale, so every child stays where it is on the page (python-pptx's own _recalculate_extents would reset the group
    to the identity transform and undo a move the user made). Does nothing when the children still fill the child space;
    goes on to the enclosing group when this one changed."""
    if grp.tag != '{http://schemas.openxmlformats.org/presentationml/2006/main}grpSp':
        return
    xf = grp.xpath('./p:grpSpPr/a:xfrm')
    kids = [r for r in (_child_rect(c) for c in grp if isinstance(c.tag, str) and c.tag.split('}')[1] in ('sp', 'grpSp', 'cxnSp', 'pic', 'graphicFrame')) if r]
    if not xf or not kids:
        return
    xf = xf[0]
    off, ext, choff, chext = (xf.find(A + t) for t in ('off', 'ext', 'chOff', 'chExt'))
    x0, y0 = min(r[0] for r in kids), min(r[1] for r in kids)
    x1, y1 = max(r[0] + r[2] for r in kids), max(r[1] + r[3] for r in kids)
    cx, cy, cw, ch = int(choff.get('x')), int(choff.get('y')), int(chext.get('cx')), int(chext.get('cy'))
    if max(abs(x0 - cx), abs(y0 - cy), abs(x1 - x0 - cw), abs(y1 - y0 - ch)) <= 12:        # 12 EMU = 0.001 pt
        return
    sx = int(ext.get('cx')) / cw if cw else 1.0
    sy = int(ext.get('cy')) / ch if ch else 1.0
    off.set('x', str(int(off.get('x')) + round((x0 - cx) * sx)))
    off.set('y', str(int(off.get('y')) + round((y0 - cy) * sy)))
    ext.set('cx', str(round((x1 - x0) * sx)))
    ext.set('cy', str(round((y1 - y0) * sy)))
    choff.set('x', str(x0))
    choff.set('y', str(y0))
    chext.set('cx', str(x1 - x0))
    chext.set('cy', str(y1 - y0))
    _refit(grp.getparent())


def _blank_app_props(path):
    """docProps/app.xml: Company and Manager emptied (PowerPoint stamps the organisation name there)."""
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == 'docProps/app.xml':
                data = re.sub(rb'<(Company|Manager)>[^<]*</\1>', rb'<\1></\1>', data)
            zo.writestr(it, data)
    os.replace(tmp, path)
